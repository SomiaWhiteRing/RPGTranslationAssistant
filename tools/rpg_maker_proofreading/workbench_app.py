from __future__ import annotations

import argparse
import csv
from dataclasses import replace as dc_replace
from pathlib import Path
from types import SimpleNamespace
import tkinter as tk
from tkinter import filedialog, messagebox, simpledialog, ttk

from openpyxl import Workbook, load_workbook

from workbench_app_v43 import RPGMakerProofreadingApp as V43App
from workbench_app_v4 import RPGMakerProofreadingApp as V4App
from workbench_core import (
    DataRecord,
    DictionaryWarning,
    EditableDictionaryRow,
    ErrorAlias,
    QAError,
    TextChange,
    analyze_dictionary,
    analyze_dunhao_usage,
    analyze_english_period_endings,
    analyze_space_structure,
    analyze_missing_translations,
    analyze_width,
    build_error_alias_changes,
    remove_translation_trailing_spaces,
    replace_dunhao_all,
    replace_dunhao_line_end,
    replace_terminal_english_period,
    sync_missing_source_leading_spaces,
    extract_database_dictionary_rows,
    is_face_message,
    is_narration_message,
    load_error_aliases,
    load_editable_dictionary,
    logical_fullwidth_units,
    normalize_newlines,
    render_rm2k3_controls,
    save_error_aliases,
    warning_is_covered_by_alias,
)

APP_TITLE = "RPG制作大师校对工具 v4.5.0"
CONFIG_NAME = "rpg_maker_proofreading_tool_config_v42.json"  # keep prior preferences


class RPGMakerProofreadingApp(V43App):
    @staticmethod
    def _app_dir() -> Path:
        return Path(__file__).resolve().parent

    @staticmethod
    def _bundle_dir() -> Path:
        return Path(__file__).resolve().parent

    def __init__(self, initial_input=None, initial_origin_dir=None, initial_translated_dir=None,
                 initial_dictionaries=None):
        self.missing_exclusions: set[str] = set()
        self.review_term_filter: str | None = None
        self.review_force_show_alias_managed = False
        super().__init__()
        self.title(APP_TITLE)
        self._initial_input = initial_input
        self._initial_origin_dir = initial_origin_dir
        self._initial_translated_dir = initial_translated_dir
        self._apply_initial_input()
        self._initial_dictionaries = initial_dictionaries or []
        self._startup_checks_pending = True
        self.after_idle(self._start_initial_checks)

    def _start_initial_checks(self):
        for value in self._initial_dictionaries:
            path = Path(value)
            if not path.is_file():
                continue
            try:
                self.dictionary_rows.extend(load_editable_dictionary(path))
            except Exception as exc:
                messagebox.showerror("启动辞典导入失败", f"{path}\n{exc}", parent=self)
        self._refresh_dictionary_tree()
        snap = self._source_snapshot()
        paths = {"json": ("json",), "txt": ("origin", "translated"), "excel": ("excel",)}[snap["kind"]]
        if not all(str(snap[key]) != "." and snap[key].exists() for key in paths):
            self.status_var.set("请选择有效的数据源；首次加载后将自动运行检查。")
            return
        self._scan_source()

    def _scan_done(self, result):
        super()._scan_done(result)
        if self._startup_checks_pending and self.all_records:
            self._startup_checks_pending = False
            checks = [
                self._do_width,
                self._do_speaker,
                self._do_quote_analysis,
                self._do_punct,
                self._refresh_symbol_catalog,
                self._do_ellipsis_scan,
                self._scan_english_period,
                self._scan_dunhao,
                self._scan_spaces,
                lambda: self._scan_alnum("case"),
                lambda: self._scan_alnum("charwidth"),
                self._do_duplicates,
                self._do_missing,
            ]
            if self.dictionary_rows:
                checks.append(self._initial_dictcheck)
            self.after_idle(lambda: self._run_initial_check(iter(checks)))

    def _initial_dictcheck(self):
        entries = self._resolve_conflicts(self.dictionary_rows)
        if entries is not None:
            self._show_dict_warnings(self._dictionary_warnings(entries))

    def _run_initial_check(self, checks):
        check = next(checks, None)
        if check is None:
            self.status_var.set("启动检查结束，请在各页面查看结果。")
            return
        try:
            check()
        except Exception as exc:
            messagebox.showerror("启动检查失败", str(exc), parent=self)
        self.after(1, lambda: self._run_initial_check(checks))

    # ------------------------------------------------------------------
    # Width analysis: Ambiguous width + adjustable panes + integrated rulers
    # ------------------------------------------------------------------
    def _build_width(self):
        super()._build_width()
        tab = self.tabs["width"]
        self.ambiguous_unit = tk.DoubleVar(value=0.5)
        self.width_text_height = tk.IntVar(value=6)
        self.width_preview_height = tk.IntVar(value=220)
        self.width_show_texts = tk.BooleanVar(value=True)
        self.width_show_preview = tk.BooleanVar(value=True)

        self.width_dual_pane = self.width_original_box.master.master
        self.width_preview_frame = self.width_canvas.master
        self.width_original_ruler = self._install_text_ruler(self.width_original_box)
        self.width_translation_ruler = self._install_text_ruler(self.width_translation_box)

        self.width_layout_ctl = ttk.LabelFrame(tab, text="宽度显示与框架高度", padding=4)
        ctl = self.width_layout_ctl
        ctl.grid(row=4, column=0, sticky="ew", pady=(4, 0))
        ttk.Label(ctl, text="Ambiguous字符单位：").pack(side="left")
        ttk.Combobox(ctl, textvariable=self.ambiguous_unit, state="readonly", values=[0.5, 1.0], width=5).pack(side="left")
        ttk.Label(ctl, text="（如 ·；实际字体仍以像素线为准）").pack(side="left", padx=(2, 10))
        ttk.Label(ctl, text="原/译文框高：").pack(side="left")
        ttk.Spinbox(ctl, from_=3, to=18, textvariable=self.width_text_height, width=5, command=self._apply_width_layout).pack(side="left")
        ttk.Label(ctl, text="预览高：").pack(side="left", padx=(8, 2))
        ttk.Spinbox(ctl, from_=120, to=700, increment=20, textvariable=self.width_preview_height, width=6, command=self._apply_width_layout).pack(side="left")
        ttk.Checkbutton(ctl, text="显示原文/译文部分", variable=self.width_show_texts, command=self._apply_width_layout).pack(side="left", padx=8)
        ttk.Checkbutton(ctl, text="显示当前字体预览", variable=self.width_show_preview, command=self._apply_width_layout).pack(side="left")
        ttk.Button(ctl, text="应用高度", command=self._apply_width_layout).pack(side="right")
        self._apply_width_layout()
        self._install_width_vertical_layout()

    def _install_text_ruler(self, text_widget: tk.Text):
        parent = text_widget.master
        # Shift existing text + scrollbars one row down once.
        for child in parent.grid_slaves():
            info = child.grid_info()
            try:
                row = int(info.get("row", 0))
            except Exception:
                row = 0
            child.grid_configure(row=row + 1)
        parent.rowconfigure(0, weight=0); parent.rowconfigure(1, weight=1)
        ruler = tk.Canvas(parent, height=26, background="#f8f8f8", highlightthickness=0)
        ruler.grid(row=0, column=0, sticky="ew")
        ruler.bind("<Configure>", lambda e, c=ruler: self._draw_text_ruler(c))
        return ruler

    def _draw_text_ruler(self, canvas: tk.Canvas):
        if not canvas.winfo_exists(): return
        canvas.delete("all")
        font = self._font_object(); charw = max(1, font.measure("汉")); h = max(24, canvas.winfo_height())
        max_units = max(30, int(canvas.winfo_width() / charw) + 2)
        for unit in range(0, max_units + 1):
            x = 4 + unit * charw
            major = unit % 5 == 0
            canvas.create_line(x, h - (14 if major else 7), x, h, fill="#777" if major else "#bbb")
            if major and unit:
                canvas.create_text(x + 2, 2, anchor="nw", text=str(unit), font=(font.actual("family"), max(7, font.actual("size") - 2)), fill="#555")
        for unit, color, label in [(float(self.face_limit.get()), "#2b6cb0", "头像"), (float(self.narr_limit.get()), "#c53030", "无头像")]:
            x = 4 + unit * charw
            canvas.create_line(x, 0, x, h, fill=color, width=2)
            canvas.create_text(x + 2, h - 14, anchor="nw", text=label, fill=color, font=(font.actual("family"), 7))

    def _refresh_all_rulers(self):
        for name in ["width_original_ruler", "width_translation_ruler", "editor_original_ruler", "editor_translation_ruler"]:
            c = getattr(self, name, None)
            if c is not None:
                try: self._draw_text_ruler(c)
                except Exception: pass

    def _apply_font(self):
        super()._apply_font()
        try: self._refresh_all_rulers()
        except Exception: pass

    def _apply_width_layout(self):
        try:
            h = max(3, int(self.width_text_height.get())); ph = max(100, int(self.width_preview_height.get()))
            self.width_original_box.configure(height=h); self.width_translation_box.configure(height=h)
            self.width_canvas.configure(height=ph)
            if self.width_show_texts.get(): self.width_dual_pane.grid()
            else: self.width_dual_pane.grid_remove()
            if self.width_show_preview.get(): self.width_preview_frame.grid()
            else: self.width_preview_frame.grid_remove()
            self._refresh_all_rulers(); self._draw_width_preview()
            if hasattr(self,"width_sash_top"): self._sync_width_grid_sizes()
        except Exception:
            pass

    def _do_width(self):
        try:
            records = self.active_records()
            unit_issues = analyze_width(records, self.face_limit.get(), self.narr_limit.get(), self.check_face.get(), self.check_narr.get(), self.ambiguous_unit.get())
            found: dict[tuple[str, int], SimpleNamespace] = {}
            for x in unit_issues:
                found[(x.record.uid, x.line_no)] = SimpleNamespace(record=x.record, line_no=x.line_no, face_type=x.face_type,
                                                                   units=x.width, limit=x.limit, pixels="", reason="全角单位超线")
            font = self._font_object(); ref_char = "汉"
            if self.check_font_pixels.get():
                for r in records:
                    if r.marker != "Message": continue
                    if is_face_message(r):
                        if not self.check_face.get(): continue
                        limit, typ = self.face_limit.get(), "有头像"
                    elif is_narration_message(r):
                        if not self.check_narr.get(): continue
                        limit, typ = self.narr_limit.get(), "无头像"
                    else: continue
                    ref_pixels = font.measure(ref_char * max(1, int(limit))) + font.measure(ref_char) * (limit % 1)
                    for n, line in enumerate(normalize_newlines(r.translated).split("\n"), 1):
                        vis = render_rm2k3_controls(line, runtime_placeholders=False)
                        pixels = font.measure(vis)
                        if pixels > ref_pixels:
                            key = (r.uid, n)
                            units = logical_fullwidth_units(vis, self.ambiguous_unit.get())
                            if key in found:
                                found[key].pixels = f"{pixels:.0f}"; found[key].reason = "全角单位＋当前字体像素超线"
                            else:
                                found[key] = SimpleNamespace(record=r, line_no=n, face_type=typ, units=units, limit=limit,
                                                             pixels=f"{pixels:.0f}", reason="当前字体像素超线")
            self.result_maps["width"] = {}; self._clear_tree(self.width_tree)
            for i, x in enumerate(found.values()):
                iid=f"w{i}"; self.result_maps["width"][iid]=x.record
                self.width_tree.insert("", "end", iid=iid, values=(x.face_type, x.record.file_key, x.line_no, f"{x.units:g}", x.pixels,
                    f"{x.limit:g}", x.reason, self._display(x.record.original), self._display(x.record.translated)))
            self.width_font_label.set(f"当前字体：{font.actual('family')} {font.actual('size')} pt")
            self.status_var.set(f"发现 {len(found)} 条宽度警告。Ambiguous字符按 {self.ambiguous_unit.get():g} 全角单位计算。")
            self._draw_width_preview(); self._refresh_all_rulers()
        except Exception as exc:
            messagebox.showerror("分析失败", str(exc), parent=self)

    # ------------------------------------------------------------------
    # Single editor: adjustable same-height layout + rulers
    # ------------------------------------------------------------------
    def _build_editor(self):
        super()._build_editor()
        tab = self.tabs["editor"]
        self.editor_text_height = tk.IntVar(value=6)
        self.editor_preview_height = tk.IntVar(value=190)
        self.editor_show_texts = tk.BooleanVar(value=True)
        self.editor_show_preview = tk.BooleanVar(value=True)
        self.editor_text_pane = self.editor_original.master.master
        self.editor_preview_frame = self.editor_width_canvas.master
        self.editor_original_ruler = self._install_text_ruler(self.editor_original)
        self.editor_translation_ruler = self._install_text_ruler(self.editor_translation)
        self.editor_layout_ctl = ttk.LabelFrame(tab, text="文本处理框架高度", padding=4)
        ctl = self.editor_layout_ctl
        ctl.grid(row=7, column=0, sticky="ew", pady=(4,0))
        ttk.Label(ctl,text="原/译文框高：").pack(side="left")
        ttk.Spinbox(ctl,from_=3,to=18,textvariable=self.editor_text_height,width=5,command=self._apply_editor_layout).pack(side="left")
        ttk.Label(ctl,text="预览高：").pack(side="left",padx=(8,2))
        ttk.Spinbox(ctl,from_=100,to=700,increment=20,textvariable=self.editor_preview_height,width=6,command=self._apply_editor_layout).pack(side="left")
        ttk.Checkbutton(ctl,text="显示原文/译文部分",variable=self.editor_show_texts,command=self._apply_editor_layout).pack(side="left",padx=8)
        ttk.Checkbutton(ctl,text="显示当前字体预览",variable=self.editor_show_preview,command=self._apply_editor_layout).pack(side="left")
        ttk.Button(ctl,text="应用高度",command=self._apply_editor_layout).pack(side="right")
        self._apply_editor_layout()
        self._install_editor_vertical_layout()

    def _apply_editor_layout(self):
        try:
            h=max(3,int(self.editor_text_height.get())); ph=max(90,int(self.editor_preview_height.get()))
            self.editor_original.configure(height=h); self.editor_translation.configure(height=h); self.editor_width_canvas.configure(height=ph)
            if self.editor_show_texts.get(): self.editor_text_pane.grid()
            else: self.editor_text_pane.grid_remove()
            if self.editor_show_preview.get(): self.editor_preview_frame.grid()
            else: self.editor_preview_frame.grid_remove()
            self._refresh_all_rulers(); self._draw_editor_width()
            if hasattr(self,"editor_sash_top"): self._sync_editor_grid_sizes()
        except Exception: pass

    # ------------------------------------------------------------------
    # Missing translation: reusable exclusion list
    # ------------------------------------------------------------------
    def _build_missing(self):
        super()._build_missing()
        tab=self.tabs["missing"]
        extra=ttk.Frame(tab); extra.grid(row=2,column=0,sticky="ew",pady=(4,0))
        ttk.Button(extra,text="将勾选项标记为“不算缺译”",command=self._missing_add_exclusions).pack(side="left")
        ttk.Button(extra,text="导入“不算缺译”文本…",command=self._missing_import_exclusions).pack(side="left",padx=4)
        ttk.Button(extra,text="导出“不算缺译”文本…",command=self._missing_export_exclusions).pack(side="left")
        ttk.Button(extra,text="清空排除表",command=self._missing_clear_exclusions).pack(side="left",padx=4)
        self.missing_exclusion_label=tk.StringVar(value="不算缺译：0 条")
        ttk.Label(extra,textvariable=self.missing_exclusion_label,foreground="#555").pack(side="left",padx=10)

    def _do_missing(self):
        try:
            selected={k for k,v in self.missing_type_vars.items() if v.get()}
            issues=[x for x in analyze_missing_translations(self.active_records(),selected) if x.record.original not in self.missing_exclusions]
            self._clear_tree(self.missing_tree); self.result_maps["missing"]={}
            for i,x in enumerate(issues):
                iid=f"miss{i}"; self.result_maps["missing"][iid]=x.record
                self.missing_tree.insert("","end",iid=iid,values=("☐",x.record.file_key,x.text_kind,x.record.marker,x.reason,
                    self._display(x.record.original),self._display(x.record.translated)))
            self.status_var.set(f"疑似日文未翻译：{len(issues)} 条；已排除“不算缺译” {len(self.missing_exclusions)} 个原文。")
        except Exception as exc: messagebox.showerror("缺译检查失败",str(exc),parent=self)

    def _missing_add_exclusions(self):
        added=0
        for iid in self._checked_iids(self.missing_tree):
            rec=self.result_maps.get("missing",{}).get(iid)
            if isinstance(rec,DataRecord) and rec.original not in self.missing_exclusions:
                self.missing_exclusions.add(rec.original); added+=1
        self.missing_exclusion_label.set(f"不算缺译：{len(self.missing_exclusions)} 条"); self._do_missing()
        self.status_var.set(f"新增 {added} 条“不算缺译”文本。")

    def _missing_import_exclusions(self):
        p=filedialog.askopenfilename(parent=self,filetypes=[("文本表","*.xlsx *.xlsm *.csv *.txt")])
        if not p:return
        try:
            values=self._read_first_column(Path(p)); before=len(self.missing_exclusions); self.missing_exclusions.update(x for x in values if x)
            self.missing_exclusion_label.set(f"不算缺译：{len(self.missing_exclusions)} 条")
            self.status_var.set(f"导入 {len(self.missing_exclusions)-before} 条“不算缺译”文本。")
        except Exception as exc: messagebox.showerror("导入失败",str(exc),parent=self)

    def _missing_export_exclusions(self):
        if not self.missing_exclusions:return
        p=filedialog.asksaveasfilename(parent=self,defaultextension=".xlsx",filetypes=[("Excel","*.xlsx"),("CSV","*.csv"),("TXT","*.txt")])
        if not p:return
        path=Path(p); values=sorted(self.missing_exclusions)
        if path.suffix.lower()==".csv":
            with path.open("w",encoding="utf-8-sig",newline="") as f:
                w=csv.writer(f);w.writerow(["不算缺译的原文"]);w.writerows([[x] for x in values])
        elif path.suffix.lower()==".txt":
            path.write_text("\n".join(values),encoding="utf-8-sig")
        else:
            wb=Workbook();ws=wb.active;ws.title="NotMissing";ws.append(["不算缺译的原文"])
            for x in values:ws.append([x])
            wb.save(path);wb.close()
        messagebox.showinfo("导出完成",str(path),parent=self)

    def _missing_clear_exclusions(self):
        if self.missing_exclusions and messagebox.askyesno("清空排除表","确定清空所有“不算缺译”文本？",parent=self):
            self.missing_exclusions.clear();self.missing_exclusion_label.set("不算缺译：0 条");self._do_missing()

    @staticmethod
    def _read_first_column(path:Path):
        if path.suffix.lower()==".txt":return [x.rstrip("\r\n") for x in path.read_text(encoding="utf-8-sig").splitlines() if x.strip()]
        if path.suffix.lower()==".csv":
            with path.open("r",encoding="utf-8-sig",newline="") as f: rows=list(csv.reader(f))
        else:
            wb=load_workbook(path,read_only=True,data_only=False)
            try: rows=[list(x) for x in wb[wb.sheetnames[0]].iter_rows(values_only=True)]
            finally:wb.close()
        if not rows:return []
        start=1 if str(rows[0][0] or "").strip() in {"不算缺译的原文","原文","Original"} else 0
        return [str(r[0] or "") for r in rows[start:] if r and str(r[0] or "").strip()]

    # ------------------------------------------------------------------
    # Dictionary page: database names (items/monsters/skills) into dictionary
    # ------------------------------------------------------------------
    def _build_dictionary(self):
        super()._build_dictionary()
        tab=self.tabs["dictionary"]
        extra=ttk.Frame(tab);extra.grid(row=3,column=0,sticky="ew",pady=(4,0))
        ttk.Button(extra,text="从当前数据库导入道具/怪物/技能名称…",command=self._import_database_names_to_dictionary).pack(side="left")
        ttk.Label(extra,text="只导入 Name，不导入 Description/UseMessage。类别：物品、怪物、技能。",foreground="#555").pack(side="left",padx=8)
        search=ttk.Frame(tab);search.grid(row=4,column=0,sticky="ew",pady=(4,0))
        self.dictionary_query=tk.StringVar()
        ttk.Label(search,text="快速查询辞典：").pack(side="left")
        ent=ttk.Entry(search,textvariable=self.dictionary_query,width=34);ent.pack(side="left",padx=3)
        ttk.Button(search,text="查询",command=self._refresh_dictionary_tree).pack(side="left")
        ttk.Button(search,text="清空",command=lambda:(self.dictionary_query.set(""),self._refresh_dictionary_tree())).pack(side="left",padx=3)
        ttk.Label(search,text="匹配原文、译文、类别和来源；定位后仍可直接编辑。",foreground="#555").pack(side="left",padx=8)
        ent.bind("<Return>",lambda e:self._refresh_dictionary_tree())
        # Extend category combobox if we can locate it.
        for widget in tab.winfo_children():
            for child in widget.winfo_children():
                if isinstance(child,ttk.Combobox):
                    try:
                        if str(child.cget("textvariable"))==str(self.dict_category):
                            vals=list(child.cget("values"));
                            for x in ["怪物","技能"]:
                                if x not in vals:vals.append(x)
                            child.configure(values=vals)
                    except Exception:pass

    def _import_database_names_to_dictionary(self):
        dlg=tk.Toplevel(self);dlg.title("导入数据库名称到辞典");dlg.transient(self);dlg.resizable(False,False)
        vals={"items":tk.BooleanVar(value=True),"monsters":tk.BooleanVar(value=True),"skills":tk.BooleanVar(value=True)}
        ttk.Label(dlg,text="选择要导入的数据库名称（仅 Name）：").pack(anchor="w",padx=12,pady=(10,4))
        ttk.Checkbutton(dlg,text="道具名 → 类别“物品”",variable=vals["items"]).pack(anchor="w",padx=20)
        ttk.Checkbutton(dlg,text="怪物名 → 类别“怪物”",variable=vals["monsters"]).pack(anchor="w",padx=20)
        ttk.Checkbutton(dlg,text="技能名 → 类别“技能”",variable=vals["skills"]).pack(anchor="w",padx=20)
        def apply():
            rows=extract_database_dictionary_rows(self.active_records(),vals["items"].get(),vals["monsters"].get(),vals["skills"].get())
            self.dictionary_rows.extend(rows);self._refresh_dictionary_tree();dlg.destroy();self.status_var.set(f"已从数据库名称加入辞典 {len(rows)} 行。")
        b=ttk.Frame(dlg);b.pack(fill="x",padx=10,pady=10);ttk.Button(b,text="导入",command=apply).pack(side="right");ttk.Button(b,text="取消",command=dlg.destroy).pack(side="right",padx=5)

    # ------------------------------------------------------------------
    # Dictionary warning/error-alias workflow v4.4
    # ------------------------------------------------------------------
    def _build_dictcheck(self):
        super()._build_dictcheck()
        tab=self.tabs["dictcheck"]
        extra=ttk.Frame(tab);extra.grid(row=3,column=0,sticky="ew",pady=(4,0))
        ttk.Button(extra,text="逐句识别：当前辞典条目",command=lambda:self._open_dictionary_review_v44("single")).pack(side="left")
        ttk.Button(extra,text="逐句识别：辞典全部条目",command=lambda:self._open_dictionary_review_v44("all")).pack(side="left",padx=4)
        ttk.Button(extra,text="导入错误译名表…",command=self._import_error_aliases).pack(side="left",padx=(12,4))
        ttk.Button(extra,text="选择并导出错误译名…",command=self._export_error_aliases_v44).pack(side="left")
        ttk.Label(extra,text="已确认错误译名会立即离开人工待确认队列；只有点击“重新检查”才重新进入。",foreground="#555").pack(side="left",padx=8)

    def _selected_main_dictionary_warning(self):
        sel=self.dictcheck_tree.selection() if hasattr(self,"dictcheck_tree") else ()
        if not sel:return None
        iid=sel[0]
        if iid.startswith("dc"):
            try:
                idx=int(iid[2:]);return self.current_dict_warnings[idx] if 0<=idx<len(self.current_dict_warnings) else None
            except Exception:return None
        return None

    def _open_dictionary_review(self):
        # Preserve old toolbar action as "all entries".
        self._open_dictionary_review_v44("all")

    def _open_dictionary_review_v44(self,mode="all"):
        if not self.current_dict_warnings and not self.error_aliases:
            messagebox.showinfo("辞典逐句确认","请先执行一次辞典匹配检查。",parent=self);return
        selected=self._selected_main_dictionary_warning() if mode=="single" else None
        if mode=="single" and selected is None:
            messagebox.showinfo("当前辞典条目","请先在上方辞典匹配结果中选择一条记录。",parent=self);return
        self.review_term_filter=selected.original_term if selected else None
        self.review_force_show_alias_managed=False
        self._mark_alias_managed_warnings()
        win=tk.Toplevel(self);self.review_window=win
        win.title("辞典匹配逐句确认 / 错误译名修正"+(f"｜仅：{self.review_term_filter}" if self.review_term_filter else "｜全部条目"));win.geometry("1550x900");win.transient(self)
        pan=ttk.PanedWindow(win,orient="horizontal");pan.pack(fill="both",expand=True,padx=6,pady=6)
        left=ttk.Frame(pan);right=ttk.Frame(pan);pan.add(left,weight=2);pan.add(right,weight=3)
        left.rowconfigure(1,weight=1);left.columnconfigure(0,weight=1)
        top=ttk.Frame(left);top.grid(row=0,column=0,sticky="ew")
        ttk.Button(top,text="重新检查（让已处理项重新进入人工列表）",command=lambda:self._review_rescan_warnings_v44(win)).pack(side="left")
        ttk.Button(top,text="跳过当前误命中",command=lambda:self._review_skip(win)).pack(side="left",padx=4)
        old_page=self._current_tree_page;self._current_tree_page="dictcheck"
        self.review_tree=self._tree(left,("status","term","expected","file"),("状态","辞典原词","正确译名","文件"),(150,220,220,280))
        self.review_tree.master.grid(row=1,column=0,sticky="nsew");self.review_tree.bind("<<TreeviewSelect>>",lambda e:self._review_show_current())
        self._current_tree_page=old_page

        right.rowconfigure(1,weight=1);right.rowconfigure(3,weight=1);right.columnconfigure(0,weight=1)
        ttk.Label(right,text="原文（可划选新辞典原词）").grid(row=0,column=0,sticky="w")
        self.review_original=tk.Text(right,height=6,wrap="word");self.review_original.grid(row=1,column=0,sticky="nsew")
        ttk.Label(right,text="译文（可划选错误译名；或划选新辞典译文）").grid(row=2,column=0,sticky="w")
        self.review_translated=tk.Text(right,height=6,wrap="word");self.review_translated.grid(row=3,column=0,sticky="nsew")
        actions=ttk.Frame(right);actions.grid(row=4,column=0,sticky="ew",pady=4)
        ttk.Button(actions,text="将译文选区加入错误译名",command=lambda:self._review_add_error_alias_v44(win)).pack(side="left")
        ttk.Button(actions,text="用原文/译文选区新增辞典项",command=lambda:self._review_add_dictionary_entry(win)).pack(side="left",padx=4)
        ttk.Label(actions,text="加入错误译名后，所有可由该错误译名自动处理的句子会立刻从上方人工列表移除。",foreground="#555").pack(side="left",padx=8)

        bottom=ttk.PanedWindow(win,orient="horizontal");bottom.pack(fill="both",expand=True,padx=6,pady=(0,6))
        af=ttk.Frame(bottom);pf=ttk.Frame(bottom);bottom.add(af,weight=2);bottom.add(pf,weight=4)
        af.rowconfigure(1,weight=1);af.columnconfigure(0,weight=1)
        ab=ttk.Frame(af);ab.grid(row=0,column=0,sticky="ew")
        ttk.Button(ab,text="删除选中错误译名",command=lambda:self._review_delete_alias(win)).pack(side="left")
        ttk.Button(ab,text="导入错误译名表…",command=lambda:self._import_error_aliases(refresh_window=win)).pack(side="left",padx=4)
        ttk.Button(ab,text="选择并导出…",command=self._export_error_aliases_v44).pack(side="left")
        old_page=self._current_tree_page;self._current_tree_page="dictcheck"
        self.alias_tree=self._tree(af,("use","wrong","correct","source_term","source"),("选择","错误译名","正确译名","对应辞典原词","来源"),(60,190,190,190,180),checkbox=True)
        self.alias_tree.master.grid(row=1,column=0,sticky="nsew")

        pf.rowconfigure(1,weight=1);pf.columnconfigure(0,weight=1)
        pb=ttk.Frame(pf);pb.grid(row=0,column=0,sticky="ew")
        ttk.Button(pb,text="全选自动替换",command=lambda:self._check_all(self.auto_tree,True)).pack(side="left")
        ttk.Button(pb,text="全不选",command=lambda:self._check_all(self.auto_tree,False)).pack(side="left",padx=3)
        ttk.Button(pb,text="应用勾选自动替换",command=lambda:self._review_apply_auto(win)).pack(side="left",padx=8)
        ttk.Label(pb,text="正确译名本身受保护；短错误词不会在正确译名内部再次扩写。",foreground="#555").pack(side="left",padx=8)
        self.auto_tree=self._tree(pf,("use","file","reason","original","translated","proposed"),("选择","文件","命中错误译名","原文","当前译文","替换预览"),(60,190,180,340,340,340),checkbox=True)
        self.auto_tree.master.grid(row=1,column=0,sticky="nsew");self.auto_tree.bind("<Double-1>",self._review_auto_double)
        self._current_tree_page=old_page
        self.review_warning_map={};self.auto_change_map={}
        self._review_refresh_warning_tree_v44();self._review_refresh_alias_tree_v44();self._review_refresh_auto_v44(win)

    def _review_active_aliases(self):
        if not self.review_term_filter:return list(self.error_aliases)
        return [x for x in self.error_aliases if x.source_term==self.review_term_filter]

    def _mark_alias_managed_warnings(self):
        aliases=self._review_active_aliases()
        if not aliases:return
        for w in self.current_dict_warnings:
            if self.review_term_filter and w.original_term!=self.review_term_filter:continue
            if warning_is_covered_by_alias(w,aliases):self.dict_review_states[w.key]="自动替换候选"

    def _review_refresh_warning_tree_v44(self):
        if not hasattr(self,"review_tree"):return
        self._clear_tree(self.review_tree);self.review_warning_map={}
        n=0
        for w in self.current_dict_warnings:
            if self.review_term_filter and w.original_term!=self.review_term_filter:continue
            state=self.dict_review_states.get(w.key,"待确认")
            if state!="待确认":continue
            iid=f"w{n}";n+=1;self.review_warning_map[iid]=w
            self.review_tree.insert("","end",iid=iid,values=(state,w.original_term,w.expected_translation,w.record.file_key))
        kids=self.review_tree.get_children("")
        if kids:self.review_tree.selection_set(kids[0]);self.review_tree.focus(kids[0]);self._review_show_current()

    def _review_add_error_alias_v44(self,win):
        w=self._review_current_warning();wrong=self._selected_text(self.review_translated).strip()
        if not w or not wrong:
            messagebox.showinfo("错误译名","请先在译文框中划选实际翻错的词。",parent=win);return
        correct=w.expected_translation.strip()
        if w.original_term=="（错误译名二次确认）":
            inferred=next((a.correct for a in sorted(self.error_aliases,key=lambda x:-len(x.wrong)) if a.wrong in wrong),"")
            if inferred:correct=inferred
        if not correct:correct=simpledialog.askstring("错误译名",f"请输入“{wrong}”应统一为的正确译名：",parent=win) or ""
        if not correct:return
        source_term = self.review_term_filter or w.original_term
        if w.original_term == "（错误译名二次确认）":
            inferred_term = next((a.source_term for a in sorted(self.error_aliases,key=lambda x:-len(x.wrong)) if a.wrong in wrong and a.source_term), "")
            source_term = self.review_term_filter or inferred_term or "（二次确认）"
        self._add_error_alias(ErrorAlias(wrong,correct,source_term,"辞典逐句确认"))
        # Immediately suppress *all* manual warnings now handled by the alias.
        aliases=self._review_active_aliases()
        covered=0
        for warning in self.current_dict_warnings:
            if self.review_term_filter and warning.original_term!=self.review_term_filter:continue
            if warning_is_covered_by_alias(warning,aliases):
                self.dict_review_states[warning.key]="自动替换候选";covered+=1
        self._review_refresh_warning_tree_v44();self._review_refresh_alias_tree_v44();self._review_refresh_auto_v44(win)
        self.status_var.set(f"已加入错误译名 {wrong} → {correct}；{covered} 条对应人工警告转入自动替换预览。")

    def _review_refresh_alias_tree_v44(self):
        if not hasattr(self,"alias_tree"):return
        self._clear_tree(self.alias_tree)
        for n,x in enumerate(sorted(self.error_aliases,key=lambda z:(z.source_term,-len(z.wrong),z.wrong))):
            self.alias_tree.insert("","end",iid=f"a{n}",values=("☐",x.wrong,x.correct,x.source_term,x.source))

    def _review_refresh_alias_tree(self):
        # Called by inherited delete/import helpers.
        if hasattr(self,"alias_tree") and "use" in self.alias_tree["columns"]:self._review_refresh_alias_tree_v44()
        else:super()._review_refresh_alias_tree()

    def _review_refresh_warning_tree(self):
        if hasattr(self,"review_tree") and self.review_tree.winfo_exists():self._review_refresh_warning_tree_v44()

    def _review_refresh_auto_v44(self,win):
        if not hasattr(self,"auto_tree"):return
        self._clear_tree(self.auto_tree)
        aliases=self._review_active_aliases()
        records=self.active_records()
        if self.review_term_filter:records=[r for r in records if self.review_term_filter in r.original]
        changes=build_error_alias_changes(records,aliases);self.auto_change_map={}
        for n,c in enumerate(changes):
            iid=f"p{n}";self.auto_change_map[iid]=c
            self.auto_tree.insert("","end",iid=iid,values=("☑",c.record.file_key,c.reason,self._display(c.record.original),self._display(c.record.translated),self._display(c.proposed)))

    def _review_refresh_auto(self,win):
        if hasattr(self,"auto_tree") and "use" in self.auto_tree["columns"]:self._review_refresh_auto_v44(win)
        else:super()._review_refresh_auto(win)

    def _review_rescan_warnings_v44(self,win):
        # Explicit recheck is the only operation that deliberately lets previously
        # alias-managed warnings re-enter the manual queue.
        self.dict_review_states={}
        self.current_dict_warnings=analyze_dictionary(self.active_records(),self.current_dict_entries,self.dict_messages_only.get())
        self._show_dict_warnings(self.current_dict_warnings);self._review_refresh_warning_tree_v44();self._review_refresh_auto_v44(win)

    def _import_error_aliases(self,refresh_window=None):
        paths=filedialog.askopenfilenames(parent=self,filetypes=[("错误译名表","*.xlsx *.xlsm *.csv")])
        if not paths:return
        added=0
        try:
            for p in paths:
                for alias in load_error_aliases(Path(p)):
                    before=len(self.error_aliases);self._add_error_alias(alias);added+=1 if len(self.error_aliases)>=before else 0
            self._mark_alias_managed_warnings()
            if refresh_window is not None:
                self._review_refresh_warning_tree_v44();self._review_refresh_alias_tree_v44();self._review_refresh_auto_v44(refresh_window)
            self.status_var.set(f"已导入/更新错误译名 {added} 项；当前共 {len(self.error_aliases)} 项。")
        except Exception as exc:messagebox.showerror("导入错误译名失败",str(exc),parent=self)

    def _export_error_aliases_v44(self):
        if not self.error_aliases:
            messagebox.showinfo("错误译名","当前没有错误译名。",parent=self);return
        top=tk.Toplevel(self);top.title("选择要导出的错误译名");top.geometry("850x620");top.transient(self)
        terms=["全部"]+sorted({x.source_term or "（无对应辞典原词）" for x in self.error_aliases})
        term=tk.StringVar(value="全部")
        bar=ttk.Frame(top);bar.pack(fill="x",padx=6,pady=6)
        ttk.Label(bar,text="对应辞典原词：").pack(side="left")
        combo=ttk.Combobox(bar,textvariable=term,state="readonly",values=terms,width=28);combo.pack(side="left")
        old_page=self._current_tree_page;self._current_tree_page="dictcheck"
        tree=self._tree(top,("use","wrong","correct","source_term","source"),("选择","错误译名","正确译名","对应辞典原词","来源"),(60,190,190,190,180),checkbox=True)
        tree.master.pack(fill="both",expand=True,padx=6,pady=(0,6));self._current_tree_page=old_page
        amap={}
        def refresh(*_):
            self._clear_tree(tree);amap.clear();target=term.get()
            for n,x in enumerate(sorted(self.error_aliases,key=lambda z:(z.source_term,-len(z.wrong),z.wrong))):
                lab=x.source_term or "（无对应辞典原词）"
                if target!="全部" and lab!=target:continue
                iid=f"e{n}";amap[iid]=x;tree.insert("","end",iid=iid,values=("☐",x.wrong,x.correct,x.source_term,x.source))
        combo.bind("<<ComboboxSelected>>",refresh);refresh()
        buttons=ttk.Frame(top);buttons.pack(fill="x",padx=6,pady=(0,6))
        ttk.Button(buttons,text="全选当前显示",command=lambda:self._check_all(tree,True)).pack(side="left")
        ttk.Button(buttons,text="全不选",command=lambda:self._check_all(tree,False)).pack(side="left",padx=4)
        def export():
            aliases=[amap[i] for i in self._checked_iids(tree) if i in amap]
            if not aliases:
                messagebox.showinfo("错误译名","请至少勾选一项。",parent=top);return
            p=filedialog.asksaveasfilename(parent=top,defaultextension=".xlsx",filetypes=[("Excel","*.xlsx"),("CSV","*.csv")])
            if p:save_error_aliases(Path(p),aliases);messagebox.showinfo("导出完成",p,parent=top)
        ttk.Button(buttons,text="导出勾选项…",command=export).pack(side="right")


    # ------------------------------------------------------------------
    # v4.5 vertical three-section layout: list/context | original+translation | font preview
    # ------------------------------------------------------------------
    def _line_pixel_height(self):
        try:
            f=self._font_object(); return max(20, int(f.metrics("linespace"))+6)
        except Exception:
            return 24

    def _make_vertical_sash(self, parent, row, start_cb, drag_cb):
        sash=ttk.Frame(parent, height=7, cursor="sb_v_double_arrow", relief="groove")
        sash.grid(row=row,column=0,sticky="ew",pady=1)
        sash.grid_propagate(False)
        sash.bind("<ButtonPress-1>",start_cb)
        sash.bind("<B1-Motion>",drag_cb)
        return sash

    def _install_width_vertical_layout(self):
        tab=self.tabs["width"]
        self.width_tree.master.grid_configure(row=1,sticky="nsew")
        self.width_dual_pane.grid_configure(row=3,sticky="nsew",pady=(2,2))
        self.width_preview_frame.grid_configure(row=5,sticky="nsew")
        self.width_layout_ctl.grid_configure(row=6,sticky="ew",pady=(4,0))
        self.width_sash_top=self._make_vertical_sash(tab,2,self._width_sash1_start,self._width_sash1_drag)
        self.width_sash_bottom=self._make_vertical_sash(tab,4,self._width_sash2_start,self._width_sash2_drag)
        for r in range(0,7): tab.rowconfigure(r,weight=0)
        tab.rowconfigure(1,weight=1,minsize=120)
        self._sync_width_grid_sizes()

    def _sync_width_grid_sizes(self):
        if not hasattr(self,"width_dual_pane"): return
        tab=self.tabs["width"]; lh=self._line_pixel_height()
        tab.rowconfigure(3,minsize=max(110,int(self.width_text_height.get())*lh+48))
        tab.rowconfigure(5,minsize=max(100,int(self.width_preview_height.get())))

    def _width_sash1_start(self,e):
        self._width_drag=(e.y_root,int(self.width_text_height.get()),int(self.width_preview_height.get()))
    def _width_sash1_drag(self,e):
        if not hasattr(self,"_width_drag"):return
        y,h,p=self._width_drag; dy=e.y_root-y; lh=self._line_pixel_height()
        self.width_text_height.set(max(3,min(18,h-round(dy/lh)))); self._apply_width_layout()
    def _width_sash2_start(self,e):
        self._width_drag=(e.y_root,int(self.width_text_height.get()),int(self.width_preview_height.get()))
    def _width_sash2_drag(self,e):
        if not hasattr(self,"_width_drag"):return
        y,h,p=self._width_drag; dy=e.y_root-y; lh=self._line_pixel_height()
        nh=max(3,min(18,h+round(dy/lh))); np=max(100,min(700,p-dy))
        self.width_text_height.set(nh); self.width_preview_height.set(np); self._apply_width_layout()

    def _install_editor_vertical_layout(self):
        tab=self.tabs["editor"]
        context_frame=self.context_tree.master.master
        # Move the transformation/operator toolbars above the three vertically resizable sections.
        extras=[]
        for child in tab.winfo_children():
            try: row=int(child.grid_info().get("row",-1))
            except Exception: row=-1
            if row in {4,5} and child not in {self.editor_text_pane,self.editor_preview_frame,self.editor_layout_ctl}:
                extras.append((row,child))
        for oldrow,child in extras:
            child.grid_configure(row=2 if oldrow==4 else 3)
        context_frame.grid_configure(row=4,sticky="nsew")
        self.editor_text_pane.grid_configure(row=6,sticky="nsew",pady=(2,2))
        self.editor_preview_frame.grid_configure(row=8,sticky="nsew")
        self.editor_layout_ctl.grid_configure(row=9,sticky="ew",pady=(4,0))
        self.editor_sash_top=self._make_vertical_sash(tab,5,self._editor_sash1_start,self._editor_sash1_drag)
        self.editor_sash_bottom=self._make_vertical_sash(tab,7,self._editor_sash2_start,self._editor_sash2_drag)
        for r in range(0,10): tab.rowconfigure(r,weight=0)
        tab.rowconfigure(4,weight=1,minsize=120)
        self._sync_editor_grid_sizes()

    def _sync_editor_grid_sizes(self):
        if not hasattr(self,"editor_text_pane"):return
        tab=self.tabs["editor"]; lh=self._line_pixel_height()
        tab.rowconfigure(6,minsize=max(110,int(self.editor_text_height.get())*lh+48))
        tab.rowconfigure(8,minsize=max(90,int(self.editor_preview_height.get())))

    def _editor_sash1_start(self,e):
        self._editor_drag=(e.y_root,int(self.editor_text_height.get()),int(self.editor_preview_height.get()))
    def _editor_sash1_drag(self,e):
        if not hasattr(self,"_editor_drag"):return
        y,h,p=self._editor_drag;dy=e.y_root-y;lh=self._line_pixel_height()
        self.editor_text_height.set(max(3,min(18,h-round(dy/lh))));self._apply_editor_layout()
    def _editor_sash2_start(self,e):
        self._editor_drag=(e.y_root,int(self.editor_text_height.get()),int(self.editor_preview_height.get()))
    def _editor_sash2_drag(self,e):
        if not hasattr(self,"_editor_drag"):return
        y,h,p=self._editor_drag;dy=e.y_root-y;lh=self._line_pixel_height()
        self.editor_text_height.set(max(3,min(18,h+round(dy/lh))));self.editor_preview_height.set(max(90,min(700,p-dy)));self._apply_editor_layout()

    # ------------------------------------------------------------------
    # Dictionary quick search
    # ------------------------------------------------------------------
    def _refresh_dictionary_tree(self):
        if not hasattr(self,"dictionary_tree"):return
        query=(self.dictionary_query.get().strip().casefold() if hasattr(self,"dictionary_query") else "")
        self._clear_tree(self.dictionary_tree)
        shown=0
        for i,r in enumerate(self.dictionary_rows):
            hay="\n".join([r.original,r.translation,r.category,r.source]).casefold()
            if query and query not in hay:continue
            self.dictionary_tree.insert("","end",iid=f"d{i}",values=(r.original,r.translation,r.category,r.source));shown+=1
        if query:self.status_var.set(f"辞典查询“{self.dictionary_query.get()}”：显示 {shown}/{len(self.dictionary_rows)} 行。")

    # ------------------------------------------------------------------
    # Punctuation workspace v4.5: three independent diagnostic tabs
    # ------------------------------------------------------------------
    def _build_punctuation_workspace(self):
        super()._build_punctuation_workspace()
        for key,title,builder in [
            ("english_period","英文句号结尾",self._build_english_period_panel),
            ("dunhao","句中顿号",self._build_dunhao_panel),
            ("spaces","空格检查",self._build_space_panel),
        ]:
            frame=ttk.Frame(self.punct_notebook,padding=6);self.punct_notebook.add(frame,text=title);self.punct_tabs[key]=frame;builder(frame)

    def _build_terminal_panel(self):
        # Use the clean sentence-terminal panel only; v4.3's combined special-format
        # button is intentionally removed in v4.5.
        V4App._build_terminal_panel(self)

    def _build_english_period_panel(self,tab):
        tab.rowconfigure(2,weight=1);tab.columnconfigure(0,weight=1)
        bar=ttk.Frame(tab);bar.grid(row=0,column=0,sticky="ew")
        ttk.Button(bar,text="检查译文英文句号结尾",command=self._scan_english_period).pack(side="left")
        ttk.Button(bar,text="全选原文不是英文句号",command=lambda:self._select_english_period("not_dot")).pack(side="left",padx=(10,3))
        ttk.Button(bar,text='全选原文是“。”句号',command=lambda:self._select_english_period("cn_dot")).pack(side="left")
        ttk.Button(bar,text="全不选",command=lambda:self._check_all(self.english_period_tree,False)).pack(side="left",padx=3)
        ttk.Button(bar,text="预览勾选：译文 . → 。",command=self._preview_english_period).pack(side="right")
        act=ttk.Frame(tab);act.grid(row=1,column=0,sticky="ew",pady=4)
        ttk.Label(act,text="原文本身以英文句号 . 结尾且译文也是 . 的记录属于正常情况，不进入列表。",foreground="#555").pack(side="left")
        ttk.Button(act,text="应用勾选预览",command=self._apply_english_period).pack(side="right")
        self.english_period_tree=self._tree(tab,("use","source_kind","source_end","file","original","translated","proposed"),
            ("选择","原文结尾类型","原文末符","文件","原文","译文","处理结果"),(60,120,80,220,430,430,430),checkbox=True)
        self.english_period_tree.master.grid(row=2,column=0,sticky="nsew");self.english_period_tree.bind("<Double-1>",lambda e:self._double_to_editor_event("english_period",self.english_period_tree,e))
        self.english_period_rows={};self.english_period_changes={}

    def _scan_english_period(self):
        rows=analyze_english_period_endings(self.active_records());self.english_period_rows={};self.english_period_changes={};self.result_maps["english_period"]={};self._clear_tree(self.english_period_tree)
        for i,row in enumerate(rows):
            iid=f"ep{i}";self.english_period_rows[iid]=row;self.result_maps["english_period"][iid]=row.record
            self.english_period_tree.insert("","end",iid=iid,values=("☐",row.source_kind,row.source_terminal or "（无）",row.record.file_key,self._display(row.record.original),self._display(row.record.translated),""))
        self.status_var.set(f"译文英文句号结尾且原文不是英文句号：{len(rows)} 条。")

    def _select_english_period(self,mode):
        for iid in self.english_period_tree.get_children(""):
            row=self.english_period_rows.get(iid);vals=list(self.english_period_tree.item(iid,"values"))
            ok=bool(row) and (mode=="not_dot" or (mode=="cn_dot" and row.source_terminal=="。"))
            vals[0]="☑" if ok else "☐";self.english_period_tree.item(iid,values=vals)

    def _preview_english_period(self):
        self.english_period_changes={}
        for iid in self._checked_iids(self.english_period_tree):
            row=self.english_period_rows.get(iid)
            if not row:continue
            proposed=replace_terminal_english_period(row.record.translated,"。")
            if proposed!=row.record.translated:self.english_period_changes[row.record.uid]=TextChange(row.record,"英文句号结尾→中文句号",proposed)
            self.english_period_tree.set(iid,"proposed",self._display(proposed) if proposed!=row.record.translated else "")
        self.status_var.set(f"英文句号结尾预览：{len(self.english_period_changes)} 条会修改。")

    def _apply_english_period(self):
        updates={uid:(c.record,c.proposed) for uid,c in self.english_period_changes.items()}
        if self._save_updates_v42(updates,"英文句号结尾"):
            self._scan_english_period()

    def _build_dunhao_panel(self,tab):
        tab.rowconfigure(2,weight=1);tab.columnconfigure(0,weight=1)
        bar=ttk.Frame(tab);bar.grid(row=0,column=0,sticky="ew")
        ttk.Button(bar,text="检查译文顿号",command=self._scan_dunhao).pack(side="left")
        for text,mode in [("全选行尾顿号","line_end"),("全选前后同字","stutter"),("全选多个顿号","multiple"),("全选其他情况","other")]:
            ttk.Button(bar,text=text,command=lambda m=mode:self._select_dunhao(m)).pack(side="left",padx=2)
        ttk.Button(bar,text="全不选",command=lambda:self._check_all(self.dunhao_tree,False)).pack(side="left",padx=3)
        act=ttk.Frame(tab);act.grid(row=1,column=0,sticky="ew",pady=4)
        ttk.Button(act,text="预览：勾选文本内所有顿号 → ，",command=lambda:self._preview_dunhao("all")).pack(side="left")
        ttk.Button(act,text="预览：只替换每行结尾顿号 → ，",command=lambda:self._preview_dunhao("line_end")).pack(side="left",padx=4)
        ttk.Button(act,text="应用勾选预览",command=self._apply_dunhao).pack(side="left",padx=8)
        ttk.Label(act,text="‘前后同字’常见于结巴；‘多个顿号’可能是排比；因此默认只检查，不自动勾选。",foreground="#555").pack(side="left",padx=10)
        self.dunhao_tree=self._tree(tab,("use","line_end","stutter","multiple","other","count","file","detail","original","translated","proposed"),
            ("选择","行尾","前后同字","多个顿号","其他","顿号数","文件","分类说明","原文","译文","处理结果"),(60,60,85,85,65,65,210,210,380,380,380),checkbox=True)
        self.dunhao_tree.master.grid(row=2,column=0,sticky="nsew");self.dunhao_tree.bind("<Double-1>",lambda e:self._double_to_editor_event("dunhao",self.dunhao_tree,e))
        self.dunhao_rows={};self.dunhao_changes={}

    def _scan_dunhao(self):
        rows=analyze_dunhao_usage(self.active_records());self.dunhao_rows={};self.dunhao_changes={};self.result_maps["dunhao"]={};self._clear_tree(self.dunhao_tree)
        for i,row in enumerate(rows):
            iid=f"dh{i}";self.dunhao_rows[iid]=row;self.result_maps["dunhao"][iid]=row.record
            self.dunhao_tree.insert("","end",iid=iid,values=("☐","是" if row.line_end_count else "否","是" if row.stutter_count else "否","是" if row.multiple else "否","是" if row.other else "否",row.total_count,row.record.file_key,row.detail,self._display(row.record.original),self._display(row.record.translated),""))
        self.status_var.set(f"译文含顿号：{len(rows)} 条文本。")

    def _select_dunhao(self,mode):
        for iid,row in self.dunhao_rows.items():
            vals=list(self.dunhao_tree.item(iid,"values"));ok={"line_end":row.line_end_count>0,"stutter":row.stutter_count>0,"multiple":row.multiple,"other":row.other}[mode]
            vals[0]="☑" if ok else "☐";self.dunhao_tree.item(iid,values=vals)

    def _preview_dunhao(self,mode):
        self.dunhao_changes={}
        for iid in self._checked_iids(self.dunhao_tree):
            row=self.dunhao_rows.get(iid)
            if not row:continue
            proposed=replace_dunhao_line_end(row.record.translated,"，") if mode=="line_end" else replace_dunhao_all(row.record.translated,"，")
            if proposed!=row.record.translated:self.dunhao_changes[row.record.uid]=TextChange(row.record,"顿号→逗号" if mode=="all" else "仅行尾顿号→逗号",proposed)
            self.dunhao_tree.set(iid,"proposed",self._display(proposed) if proposed!=row.record.translated else "")
        self.status_var.set(f"顿号替换预览：{len(self.dunhao_changes)} 条会修改。")

    def _apply_dunhao(self):
        updates={uid:(c.record,c.proposed) for uid,c in self.dunhao_changes.items()}
        if self._save_updates_v42(updates,"顿号替换"):self._scan_dunhao()

    def _build_space_panel(self,tab):
        tab.rowconfigure(2,weight=1);tab.columnconfigure(0,weight=1)
        bar=ttk.Frame(tab);bar.grid(row=0,column=0,sticky="ew")
        self.space_preserve_source_trailing=tk.BooleanVar(value=True)
        ttk.Button(bar,text="检查空格结构",command=lambda:self._scan_spaces(False)).pack(side="left")
        ttk.Button(bar,text="查看所有原文/译文带空格文本",command=lambda:self._scan_spaces(True)).pack(side="left",padx=4)
        ttk.Checkbutton(bar,text="原文同行也有行尾空格时，不删除译文行尾空格",variable=self.space_preserve_source_trailing).pack(side="left",padx=10)
        act=ttk.Frame(tab);act.grid(row=1,column=0,sticky="ew",pady=4)
        for text,cat in [("全选原文行首空格","原文行首空格"),("全选译文行尾空格","译文行尾空格"),("全选中间空格不一致","文本中间空格不一致")]:
            ttk.Button(act,text=text,command=lambda c=cat:self._select_space_category(c)).pack(side="left",padx=2)
        ttk.Button(act,text="全不选",command=lambda:self._check_all(self.space_tree,False)).pack(side="left",padx=3)
        ttk.Button(act,text="预览补齐原文行首空格",command=lambda:self._preview_spaces("leading")).pack(side="left",padx=(10,3))
        ttk.Button(act,text="预览删除译文行尾空格",command=lambda:self._preview_spaces("trailing")).pack(side="left")
        ttk.Button(act,text="应用预览",command=self._apply_spaces).pack(side="right")
        self.space_tree=self._tree(tab,("use","category","lines","auto","file","detail","original","translated","proposed"),
            ("选择","类型","行号","可自动","文件","说明","原文","译文","处理结果"),(60,150,100,80,210,300,390,390,390),checkbox=True)
        self.space_tree.master.grid(row=2,column=0,sticky="nsew");self.space_tree.bind("<Double-1>",lambda e:self._double_to_editor_event("spaces",self.space_tree,e))
        self.space_rows={};self.space_changes={}

    def _scan_spaces(self,include_all=False):
        rows=analyze_space_structure(self.active_records(),include_all_with_spaces=include_all,preserve_source_trailing=self.space_preserve_source_trailing.get())
        self.space_rows={};self.space_changes={};self.result_maps["spaces"]={};self._clear_tree(self.space_tree)
        for i,row in enumerate(rows):
            iid=f"spc{i}";self.space_rows[iid]=row;self.result_maps["spaces"][iid]=row.record
            self.space_tree.insert("","end",iid=iid,values=("☐",row.category,",".join(map(str,row.lines)),"是" if row.proposed is not None else "否",row.record.file_key,row.detail,self._display(row.record.original),self._display(row.record.translated),""))
        self.status_var.set(f"空格检查：{len(rows)} 项。")

    def _select_space_category(self,category):
        for iid,row in self.space_rows.items():
            vals=list(self.space_tree.item(iid,"values"));vals[0]="☑" if row.category==category else "☐";self.space_tree.item(iid,values=vals)

    def _preview_spaces(self,mode):
        self.space_changes={}
        selected_by_uid={}
        for iid in self._checked_iids(self.space_tree):
            row=self.space_rows.get(iid)
            if row:selected_by_uid.setdefault(row.record.uid,row.record)
        for uid,rec in selected_by_uid.items():
            text=rec.translated
            temp=rec
            if mode=="leading":
                proposed=sync_missing_source_leading_spaces(temp)
            else:
                proposed=remove_translation_trailing_spaces(temp,self.space_preserve_source_trailing.get())
            if proposed is not None and proposed!=text:self.space_changes[uid]=TextChange(rec,"补齐原文行首空格" if mode=="leading" else "删除译文行尾空格",proposed)
        for iid,row in self.space_rows.items():
            c=self.space_changes.get(row.record.uid);self.space_tree.set(iid,"proposed",self._display(c.proposed) if c else "")
        self.status_var.set(f"空格处理预览：{len(self.space_changes)} 条会修改。")

    def _apply_spaces(self):
        updates={uid:(c.record,c.proposed) for uid,c in self.space_changes.items()}
        if self._save_updates_v42(updates,"空格处理"):self._scan_spaces(False)

    # ------------------------------------------------------------------
    # Keep inherited speaker->aliases flow, but v4.4 review uses the safe engine.
    # ------------------------------------------------------------------


def main():
    parser = argparse.ArgumentParser(description=APP_TITLE)
    parser.add_argument("--initial-input")
    parser.add_argument("--initial-origin-dir")
    parser.add_argument("--initial-translated-dir")
    parser.add_argument("--initial-dictionary", action="append", default=[],
                        help="可选辞典文件，可重复传入；不存在时跳过")
    args = parser.parse_args()
    app = RPGMakerProofreadingApp(
        initial_input=args.initial_input,
        initial_origin_dir=args.initial_origin_dir,
        initial_translated_dir=args.initial_translated_dir,
        initial_dictionaries=args.initial_dictionary,
    )
    app.mainloop()


if __name__=="__main__":main()
