from __future__ import annotations

import csv
import json
import tkinter as tk
from pathlib import Path
from tkinter import filedialog, messagebox, ttk

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from workbench_app_v41 import RPGMakerProofreadingApp as V41App
from workbench_core import (
    BATCH_QA_CHECK_LABELS,
    BATCH_QA_CONTENT_KEYS,
    BATCH_QA_FORMAT_KEYS,
    CONFIDENCE_RANK,
    DataRecord,
    DuplicateTextGroup,
    FormatCheckIssue,
    QAError,
    TextChange,
    VersionDiffRow,
    analyze_ellipsis_occurrences,
    analyze_format_checks,
    analyze_missing_translations,
    analyze_punctuation,
    analyze_quote_combined,
    analyze_width,
    apply_translation_updates,
    backup_source_snapshot,
    build_duplicate_sync_changes,
    build_speaker_sync_change,
    case_format_matches,
    case_profile,
    collect_batch_qa_report_issues,
    compare_translation_versions,
    is_face_message,
    is_narration_message,
    load_config,
    load_excel_records,
    load_json_records,
    load_txt_records,
    normalize_newlines,
    render_rm2k3_controls,
    save_config,
    strip_control_codes,
    width_format_matches,
    width_profile,
)

APP_TITLE = "RPG制作大师校对工具 v4.2.1"
CONFIG_NAME = "rpg_maker_proofreading_tool_config_v42.json"

PAGE_LABELS_V42 = {
    "search": "文本查询", "width": "检查文本宽度", "editor": "单独文本处理", "speaker": "筛选说话人",
    "punctuation": "标点符号处理", "alnum": "英语数字格式", "duplicates": "重复文本检查",
    "dictionary": "生成辞典", "dictcheck": "匹配辞典检查", "database": "数据库直接翻译",
    "missing": "缺译快速检查", "versions": "版本译文差异",
}
PAGE_KEYS_V42 = {v: k for k, v in PAGE_LABELS_V42.items()}


class RPGMakerProofreadingApp(V41App):
    def __init__(self):
        self._startup_backup_signatures: set[str] = set()
        self._save_operation_count = 0
        self._last_auto_backup: Path | None = None
        self.speaker_sync_changes: dict[str, TextChange] = {}
        self.duplicate_targets: dict[str, str] = {}
        super().__init__()
        self.title(APP_TITLE)
        self.config_path = self._app_dir() / CONFIG_NAME
        # v4.1 loads its settings while constructing; reload using v4.2 config now.
        self._load_settings()
        self._update_backup_status()

    # ------------------------------------------------------------------
    # UI structure / page names
    # ------------------------------------------------------------------
    def _build_ui(self):
        super()._build_ui()
        for key, title, builder in [
            ("missing", "12. 缺译快速检查", self._build_missing),
            ("versions", "13. 版本译文差异", self._build_versions),
        ]:
            if key in self.tabs:
                continue
            frame = ttk.Frame(self.notebook, padding=8)
            self.notebook.add(frame, text=title); self.tabs[key] = frame
            self._current_tree_page = key; builder(); self._current_tree_page = ""
        self._refresh_column_page_choices()

    def _refresh_column_page_choices(self):
        if not hasattr(self, "column_page_combo"):
            return
        pages = [p for p in self.page_trees if p != "settings"]
        labels = [PAGE_LABELS_V42.get(p, p) for p in pages]
        self.column_page_combo.configure(values=labels)
        if labels and self.column_page.get() not in labels:
            self.column_page.set(labels[0])
        self._refresh_column_config_tree()

    def _column_page_key(self):
        return PAGE_KEYS_V42.get(self.column_page.get(), self.column_page.get())

    # ------------------------------------------------------------------
    # Backup policy + home-page batch report
    # ------------------------------------------------------------------
    def _annotate_dynamic_rows(self):
        super()._annotate_dynamic_rows()
        self.backup_every_n = tk.IntVar(value=10)
        self.startup_backup_enabled = tk.BooleanVar(value=True)
        self.backup_status = tk.StringVar(value="尚未备份")
        try:
            tab = self.tabs["settings"]
            src = next(w for w in tab.winfo_children() if isinstance(w, ttk.LabelFrame) and "数据源" in str(w.cget("text")))
            row = ttk.Frame(src); row.grid(row=10, column=0, columnspan=4, sticky="ew", pady=(5, 0))
            ttk.Checkbutton(row, text="首次检查数据源时自动备份", variable=self.startup_backup_enabled).pack(side="left")
            ttk.Label(row, text="每").pack(side="left", padx=(12, 2))
            ttk.Spinbox(row, from_=0, to=9999, textvariable=self.backup_every_n, width=6).pack(side="left")
            ttk.Label(row, text="次保存自动备份（0=关闭）").pack(side="left", padx=(2, 8))
            ttk.Button(row, text="立即手动备份", command=self._manual_backup).pack(side="left")
            ttk.Label(row, textvariable=self.backup_status, foreground="#555").pack(side="left", padx=10)
            row2 = ttk.Frame(src); row2.grid(row=11, column=0, columnspan=4, sticky="ew", pady=(4, 0))
            ttk.Button(row2, text="批量输出格式检查结果…", command=self._export_format_report).pack(side="left")
            ttk.Label(row2, text="导出前可选择检查项目，包括格式问题、说话人译名不统一、重复文本译文不一致等。", foreground="#555").pack(side="left", padx=10)
        except Exception:
            pass

    def _source_signature(self) -> str:
        kind = self.source_type.get()
        if kind == "json": return "json|" + str(Path(self.json_path.get()).resolve())
        if kind == "excel": return "excel|" + str(Path(self.excel_dir.get()).resolve())
        return "txt|" + str(Path(self.origin_dir.get()).resolve()) + "|" + str(Path(self.translated_dir.get()).resolve())

    def _scan_done(self, result):
        super()._scan_done(result)
        if not getattr(self, "startup_backup_enabled", None) or not self.startup_backup_enabled.get():
            return
        sig = self._source_signature()
        if sig in self._startup_backup_signatures:
            return
        try:
            root = backup_source_snapshot(self.all_records, "Startup")
            self._startup_backup_signatures.add(sig)
            self.backup_status.set(f"启动备份：{root.name}")
        except Exception as exc:
            messagebox.showwarning("首次备份失败", str(exc), parent=self)

    def _manual_backup(self):
        if not self.all_records:
            messagebox.showinfo("手动备份", "请先检查数据源。", parent=self); return
        try:
            root = backup_source_snapshot(self.all_records, "Manual")
            self.backup_status.set(f"手动备份：{root.name}")
            messagebox.showinfo("手动备份完成", str(root), parent=self)
        except Exception as exc:
            messagebox.showerror("手动备份失败", str(exc), parent=self)

    def _update_backup_status(self):
        if not hasattr(self, "backup_status"): return
        n = self.backup_every_n.get() if hasattr(self, "backup_every_n") else 0
        suffix = f"；每 {n} 次保存自动备份" if n else "；定期自动备份关闭"
        if self._last_auto_backup:
            self.backup_status.set(f"保存次数 {self._save_operation_count}；最近自动备份 {self._last_auto_backup.name}{suffix}")
        elif self._save_operation_count:
            self.backup_status.set(f"保存次数 {self._save_operation_count}{suffix}")

    def _save_updates_v42(self, updates, title: str):
        if not updates:
            messagebox.showinfo(title, "没有可保存的修改。", parent=self); return False
        try:
            count, _ = apply_translation_updates(updates, create_backup=False)
            self._save_operation_count += 1
            self._refresh_after_write()
            auto = None
            n = max(0, int(self.backup_every_n.get())) if hasattr(self, "backup_every_n") else 0
            if n and self._save_operation_count % n == 0:
                auto = backup_source_snapshot(self.all_records, f"Auto_{self._save_operation_count}")
                self._last_auto_backup = auto
            self._update_backup_status()
            msg = f"已保存 {count} 条。\n本次未创建逐次备份。"
            if auto: msg += f"\n已完成第 {self._save_operation_count} 次保存后的自动备份：\n{auto}"
            messagebox.showinfo(title, msg, parent=self)
            return True
        except Exception as exc:
            messagebox.showerror(title + "失败", str(exc), parent=self); return False

    def _apply_updates_dialog(self, updates, title):
        self._save_updates_v42(updates, title)

    def _editor_save(self):
        if getattr(self, "editor_controls_hidden", False):
            messagebox.showwarning("不能保存", "隐藏 RPG Maker 操作符时不能保存，请先显示操作符。", parent=self); return
        if not self.current_record: return
        proposed = self.editor_translation.get("1.0", "end-1c")
        if proposed == self.current_record.translated: return
        uid = self.current_record.uid
        if self._save_updates_v42({uid: (self.current_record, proposed)}, "保存完成"):
            new = self.record_by_uid.get(uid)
            if new: self._show_editor_record(new)

    def _load_settings(self):
        super()._load_settings()
        if hasattr(self, "backup_every_n"):
            cfg = load_config(self.config_path)
            self.backup_every_n.set(int(cfg.get("backup_every_n", 10)))
            self.startup_backup_enabled.set(bool(cfg.get("startup_backup_enabled", True)))

    def _save_settings(self):
        super()._save_settings()
        cfg = load_config(self.config_path)
        if hasattr(self, "backup_every_n"):
            cfg["backup_every_n"] = int(self.backup_every_n.get())
            cfg["startup_backup_enabled"] = bool(self.startup_backup_enabled.get())
        save_config(self.config_path, cfg)

    # ------------------------------------------------------------------
    # Width page: direct edit + dual source/translation boxes + line preview
    # ------------------------------------------------------------------
    def _build_width(self):
        super()._build_width()
        tab = self.tabs["width"]
        # Move the existing line-preview frame down and put editable dual boxes above it.
        preview = self.width_canvas.master
        preview.grid_configure(row=3)
        tab.rowconfigure(1, weight=3); tab.rowconfigure(2, weight=1); tab.rowconfigure(3, weight=2)
        dual = ttk.PanedWindow(tab, orient="horizontal"); dual.grid(row=2, column=0, sticky="nsew", pady=(5, 0))
        of = ttk.LabelFrame(dual, text="原文（只读）", padding=3); tf = ttk.LabelFrame(dual, text="译文（可直接修改）", padding=3)
        dual.add(of, weight=1); dual.add(tf, weight=1)
        of.rowconfigure(0, weight=1); of.columnconfigure(0, weight=1)
        tf.rowconfigure(0, weight=1); tf.columnconfigure(0, weight=1)
        self.width_original_box = self._text(of, 6, True); self.width_translation_box = self._text(tf, 6, False)
        btn = ttk.Frame(tf); btn.grid(row=1, column=0, sticky="ew")
        ttk.Button(btn, text="保存当前译文", command=self._width_save).pack(side="right")
        ttk.Button(btn, text="恢复扫描时译文", command=self._width_reload_boxes).pack(side="right", padx=4)
        self.width_current_record: DataRecord | None = None
        self.width_tree.bind("<<TreeviewSelect>>", lambda e: self._width_selection_to_boxes(), add=True)
        self.width_translation_box.bind("<KeyRelease>", lambda e: self._draw_width_preview(), add=True)

    def _width_selection_to_boxes(self):
        sel = self.width_tree.selection()
        if not sel: return
        rec = self.result_maps.get("width", {}).get(sel[0])
        if not isinstance(rec, DataRecord): return
        self.width_current_record = rec
        self._set_text(self.width_original_box, rec.original, True)
        self._set_text(self.width_translation_box, rec.translated, False)
        self._draw_width_preview()

    def _width_reload_boxes(self):
        if self.width_current_record:
            self._set_text(self.width_translation_box, self.width_current_record.translated, False); self._draw_width_preview()

    def _width_save(self):
        rec = self.width_current_record
        if not rec: return
        proposed = self.width_translation_box.get("1.0", "end-1c")
        if proposed == rec.translated: return
        uid = rec.uid
        if self._save_updates_v42({uid: (rec, proposed)}, "宽度页保存"):
            self.width_current_record = self.record_by_uid.get(uid)
            self._do_width()

    def _draw_width_preview(self):
        if not hasattr(self, "width_canvas"): return
        canvas = self.width_canvas; canvas.delete("all")
        font = self._font_object(); family = font.actual("family"); size = font.actual("size")
        rec = getattr(self, "width_current_record", None)
        if not isinstance(rec, DataRecord):
            sel = self.width_tree.selection() if hasattr(self, "width_tree") else ()
            rec = self.result_maps.get("width", {}).get(sel[0]) if sel else self.current_record
        if not isinstance(rec, DataRecord):
            canvas.create_text(15, 15, anchor="nw", text=f"当前字体：{family} {size} pt\n选择一条记录后显示原文/译文及两类参考线。", font=(family, size)); return
        original = self.width_original_box.get("1.0", "end-1c") if hasattr(self, "width_original_box") else rec.original
        translated = self.width_translation_box.get("1.0", "end-1c") if hasattr(self, "width_translation_box") else rec.translated
        self._draw_dual_width_canvas(canvas, original, translated, rec)

    def _draw_dual_width_canvas(self, canvas, original: str, translated: str, rec: DataRecord):
        font = self._font_object(); family = font.actual("family"); size = font.actual("size")
        width = max(900, canvas.winfo_width()); half = width // 2
        height = max(220, canvas.winfo_height())
        canvas.configure(scrollregion=(0, 0, width, max(height, 320)))
        charw = font.measure("汉")
        face_limit = float(self.face_limit.get()); narr_limit = float(self.narr_limit.get())
        for base, title, text in [(15, "原文", original), (half + 15, "译文", translated)]:
            canvas.create_text(base, 8, anchor="nw", text=f"{title}｜当前字体：{family} {size} pt", font=(family, size, "bold"))
            fx = base + charw * face_limit; nx = base + charw * narr_limit
            canvas.create_line(fx, 34, fx, height - 10, fill="#2b6cb0", width=2)
            canvas.create_line(nx, 34, nx, height - 10, fill="#c53030", width=2)
            y = 40
            for line in normalize_newlines(text).split("\n")[:8]:
                canvas.create_text(base, y, anchor="nw", text=render_rm2k3_controls(line), font=(family, size)); y += max(22, size + 10)
        canvas.create_line(half, 0, half, height, fill="#aaa")

    # ------------------------------------------------------------------
    # Single editor: same line preview as width page
    # ------------------------------------------------------------------
    def _build_editor(self):
        super()._build_editor()
        self.editor_original.configure(height=6); self.editor_translation.configure(height=6)
        tab = self.tabs["editor"]
        frame = ttk.LabelFrame(tab, text="当前字体宽度参考：蓝线=有头像；红线=无头像", padding=3)
        frame.grid(row=6, column=0, sticky="ew", pady=(4, 0))
        self.editor_width_canvas = tk.Canvas(frame, background="white", height=190)
        self.editor_width_canvas.pack(fill="x", expand=True)
        self.editor_translation.bind("<KeyRelease>", lambda e: self._draw_editor_width(), add=True)

    def _show_editor_record(self, record):
        super()._show_editor_record(record)
        self.after_idle(self._draw_editor_width)

    def _draw_editor_width(self):
        if not hasattr(self, "editor_width_canvas") or not self.current_record: return
        original = self.editor_original.get("1.0", "end-1c")
        translated = self.editor_translation.get("1.0", "end-1c")
        c = self.editor_width_canvas; c.delete("all")
        self._draw_dual_width_canvas(c, original, translated, self.current_record)

    # ------------------------------------------------------------------
    # Speaker: fix double-click + safe synchronization preview/apply
    # ------------------------------------------------------------------
    def _build_speaker(self):
        super()._build_speaker()
        self.speaker_tree.bind("<Double-1>", self._speaker_group_double, add=True)
        self.speaker_example_tree.bind("<Double-1>", self._speaker_example_double)
        cols = tuple(self.speaker_example_tree["columns"])
        if "proposed" not in cols:
            self.speaker_example_tree.configure(columns=cols + ("proposed",))
            self.speaker_example_tree.heading("proposed", text="说话人同步预览")
            self.speaker_example_tree.column("proposed", width=420, minwidth=0, anchor="w")
            meta = self.tree_registry.get(self.speaker_example_tree)
            if meta:
                meta["columns"] = cols + ("proposed",); meta["headings"]["proposed"] = "说话人同步预览"; meta["widths"]["proposed"] = 420
        tab = self.tabs["speaker"]
        sync = ttk.Frame(tab); sync.grid(row=4, column=0, sticky="ew", pady=(4, 0))
        ttk.Button(sync, text="预览同步当前说话人（仅稳定项）", command=self._preview_speaker_sync).pack(side="left")
        ttk.Button(sync, text="应用稳定预览", command=self._apply_speaker_sync).pack(side="left", padx=5)
        ttk.Label(sync, text="置信度低于“较确定”或无法稳定识别译文说话人位置的例句不会修改。", foreground="#555").pack(side="left", padx=10)

    def _speaker_group_double(self, event):
        iid = self.speaker_tree.identify_row(event.y)
        group = self.result_maps.get("speaker", {}).get(iid)
        if group and getattr(group, "records", None):
            self._send_to_editor(group.records[0]); return "break"

    def _speaker_example_double(self, event):
        iid = self.speaker_example_tree.identify_row(event.y)
        rec = self.result_maps.get("speaker_examples", {}).get(iid)
        if isinstance(rec, DataRecord): self._send_to_editor(rec); return "break"

    def _show_speaker_examples(self):
        super()._show_speaker_examples()
        for iid in self.speaker_example_tree.get_children(""):
            rec = self.result_maps.get("speaker_examples", {}).get(iid)
            change = self.speaker_sync_changes.get(rec.uid) if isinstance(rec, DataRecord) else None
            self.speaker_example_tree.set(iid, "proposed", self._display(change.proposed) if change else "")

    def _preview_speaker_sync(self):
        sel = self.speaker_tree.selection()
        if not sel: return
        group = self.result_maps.get("speaker", {}).get(sel[0])
        target = self.speaker_choice.get().strip()
        if not group or not target: return
        self.speaker_sync_changes = {}
        skipped = 0
        for rec in group.records:
            change = build_speaker_sync_change(rec, target, self.speaker_mode.get())
            if change: self.speaker_sync_changes[rec.uid] = change
            else: skipped += 1
        self._show_speaker_examples()
        self.status_var.set(f"说话人同步预览：稳定可修改 {len(self.speaker_sync_changes)} 条；跳过不稳定/无需修改 {skipped} 条。")

    def _apply_speaker_sync(self):
        updates = {uid: (c.record, c.proposed) for uid, c in self.speaker_sync_changes.items()}
        if self._save_updates_v42(updates, "说话人同步"):
            self.speaker_sync_changes.clear(); self._do_speaker()

    # ------------------------------------------------------------------
    # Punctuation: English-period / internal dunhao / space check
    # ------------------------------------------------------------------
    def _build_terminal_panel(self):
        super()._build_terminal_panel()
        tab = self.punct_tabs["terminal"]
        extra = ttk.Frame(tab); extra.grid(row=3, column=0, sticky="ew", pady=(4, 0))
        ttk.Button(extra, text="检查英文句号结尾／句中顿号／空格", command=self._do_special_format_scan).pack(side="left")
        ttk.Label(extra, text="这里只生成手工检查警告，不自动替换。", foreground="#555").pack(side="left", padx=8)

    def _do_special_format_scan(self):
        issues = analyze_format_checks(self.active_records())
        self._clear_tree(self.terminal_tree); self.result_maps["terminal"] = {}
        for i, issue in enumerate(issues):
            iid = f"fmt{i}"; self.result_maps["terminal"][iid] = issue.record
            self.terminal_tree.insert("", "end", iid=iid, values=("☐", issue.record.file_key, f"{issue.side}｜{issue.issue_type}：{issue.detail}",
                self._display(issue.record.original), self._display(issue.record.translated), "（手动检查）"))
        self.status_var.set(f"格式专项警告 {len(issues)} 条。")

    # ------------------------------------------------------------------
    # Duplicate translation synchronization
    # ------------------------------------------------------------------
    def _build_duplicates(self):
        super()._build_duplicates()
        cols = tuple(self.duplicates_tree["columns"])
        if "target" not in cols:
            self.duplicates_tree.configure(columns=cols + ("target",))
            self.duplicates_tree.heading("target", text="同步目标译文")
            self.duplicates_tree.column("target", width=420, minwidth=0, anchor="w")
            meta = self.tree_registry.get(self.duplicates_tree)
            if meta:
                meta["columns"] = cols + ("target",); meta["headings"]["target"] = "同步目标译文"; meta["widths"]["target"] = 420
        tab = self.tabs["duplicates"]
        sync = ttk.Frame(tab); sync.grid(row=3, column=0, sticky="ew", pady=(4, 0))
        self.duplicate_choice = tk.StringVar()
        ttk.Label(sync, text="当前组同步成：").pack(side="left")
        self.duplicate_choice_combo = ttk.Combobox(sync, textvariable=self.duplicate_choice, state="readonly", width=50)
        self.duplicate_choice_combo.pack(side="left")
        self.duplicate_choice_combo.bind("<<ComboboxSelected>>", lambda e: self._set_duplicate_target())
        ttk.Button(sync, text="设为当前组同步目标", command=self._set_duplicate_target).pack(side="left", padx=4)
        ttk.Button(sync, text="应用勾选组的已设目标", command=self._apply_duplicate_sync).pack(side="left", padx=4)
        self.duplicates_tree.bind("<<TreeviewSelect>>", lambda e: self._duplicate_selection_changed(), add=True)

    def _duplicate_selection_changed(self):
        sel = self.duplicates_tree.selection()
        if not sel: return
        iid = sel[0]; group = self.result_maps.get("duplicates", {}).get(iid)
        if not isinstance(group, DuplicateTextGroup): return
        vals = [t for t, _n in group.translations]
        self.duplicate_choice_combo.configure(values=vals)
        self.duplicate_choice.set(self.duplicate_targets.get(iid, vals[0] if vals else ""))

    def _set_duplicate_target(self):
        sel = self.duplicates_tree.selection()
        if not sel: return
        iid = sel[0]; target = self.duplicate_choice.get()
        if not target: return
        self.duplicate_targets[iid] = target; self.duplicates_tree.set(iid, "target", self._display(target))

    def _apply_duplicate_sync(self):
        updates = {}
        missing = []
        for iid in self._checked_iids(self.duplicates_tree):
            group = self.result_maps.get("duplicates", {}).get(iid)
            if not isinstance(group, DuplicateTextGroup): continue
            target = self.duplicate_targets.get(iid)
            if not target:
                missing.append(iid); continue
            for change in build_duplicate_sync_changes(group, target):
                updates[change.record.uid] = (change.record, change.proposed)
        if missing:
            messagebox.showwarning("重复文本同步", f"有 {len(missing)} 个勾选组尚未设置“同步目标译文”，这些组不会修改。", parent=self)
        if self._save_updates_v42(updates, "重复文本同步"):
            self.duplicate_targets.clear(); self._do_duplicates()

    # ------------------------------------------------------------------
    # Missing translation / original==translation check
    # ------------------------------------------------------------------
    def _build_missing(self):
        tab = self.tabs["missing"]; tab.rowconfigure(1, weight=1); tab.columnconfigure(0, weight=1)
        bar = ttk.Frame(tab); bar.grid(row=0, column=0, sticky="ew", pady=(0, 5))
        self.missing_filter = tk.StringVar(value="全部")
        ttk.Label(bar, text="检查范围：").pack(side="left")
        ttk.Combobox(bar, textvariable=self.missing_filter, state="readonly", values=["全部", "全汉字", "全英文"], width=12).pack(side="left")
        ttk.Button(bar, text="快速检查", command=self._do_missing).pack(side="left", padx=5)
        ttk.Button(bar, text="全选当前可见行", command=lambda: self._check_all(self.missing_tree, True)).pack(side="left", padx=(10, 3))
        ttk.Button(bar, text="全不选", command=lambda: self._check_all(self.missing_tree, False)).pack(side="left")
        ttk.Button(bar, text="打开译文文件", command=lambda: self._open_selected_result("missing", "translated")).pack(side="right")
        self.missing_tree = self._tree(tab, ("use", "file", "kind", "marker", "reason", "original", "translated"),
            ("选择", "文件", "原文类型", "标记", "说明", "原文", "译文"), (60, 240, 100, 100, 220, 500, 500), checkbox=True)
        self.missing_tree.master.grid(row=1, column=0, sticky="nsew")
        self.missing_tree.bind("<Double-1>", lambda e: self._double_to_editor_event("missing", self.missing_tree, e))

    def _do_missing(self):
        try:
            issues = analyze_missing_translations(self.active_records(), self.missing_filter.get())
            self._clear_tree(self.missing_tree); self.result_maps["missing"] = {}
            for i, x in enumerate(issues):
                iid = f"miss{i}"; self.result_maps["missing"][iid] = x.record
                self.missing_tree.insert("", "end", iid=iid, values=("☐", x.record.file_key, x.text_kind, x.record.marker, x.reason,
                    self._display(x.record.original), self._display(x.record.translated)))
            self.status_var.set(f"疑似缺译/原译相同 {len(issues)} 条（纯标点和空文本已排除）。")
        except Exception as exc: messagebox.showerror("缺译检查失败", str(exc), parent=self)

    # ------------------------------------------------------------------
    # Version translation comparison
    # ------------------------------------------------------------------
    def _build_versions(self):
        tab = self.tabs["versions"]; tab.rowconfigure(3, weight=1); tab.columnconfigure(1, weight=1)
        self.version_source_type = tk.StringVar(value="json")
        self.version_json = tk.StringVar(); self.version_origin = tk.StringVar(); self.version_translated = tk.StringVar(); self.version_excel = tk.StringVar()
        top = ttk.Frame(tab); top.grid(row=0, column=0, columnspan=2, sticky="ew")
        ttk.Label(top, text="对照旧版本格式：").pack(side="left")
        for text, val in [("JSON", "json"), ("TXT", "txt"), ("Excel", "excel")]:
            ttk.Radiobutton(top, text=text, value=val, variable=self.version_source_type, command=self._version_visibility).pack(side="left", padx=4)
        self.version_rows = {}
        self.version_rows["json"] = self._version_path_row(tab, 1, "旧版本 JSON", self.version_json, False)
        self.version_rows["origin"] = self._version_path_row(tab, 1, "旧版本原文 TXT（可留空=当前原文）", self.version_origin, True)
        self.version_rows["translated"] = self._version_path_row(tab, 2, "旧版本译文 TXT", self.version_translated, True)
        self.version_rows["excel"] = self._version_path_row(tab, 1, "旧版本 Excel 文件夹", self.version_excel, True)
        actions = ttk.Frame(tab); actions.grid(row=2, column=0, columnspan=2, sticky="ew", pady=5)
        ttk.Button(actions, text="比较译文修改", command=self._do_version_compare).pack(side="left")
        ttk.Button(actions, text="导出差异报告…", command=self._export_version_report).pack(side="left", padx=5)
        ttk.Label(actions, text="仅比较原文完全一致的对应记录；原文变化的记录不纳入译文修改报告。", foreground="#555").pack(side="left", padx=10)
        self.version_tree = self._tree(tab, ("file", "original", "old", "new", "old_location", "new_location"),
            ("文件", "原文", "旧译文", "新译文", "旧版位置", "新版位置"), (220, 450, 450, 450, 380, 380))
        self.version_tree.master.grid(row=3, column=0, columnspan=2, sticky="nsew")
        self.version_tree.bind("<Double-1>", lambda e: self._double_to_editor_event("versions", self.version_tree, e))
        self.version_results: list[VersionDiffRow] = []
        self._version_visibility()

    def _version_path_row(self, parent, row, label, var, directory: bool):
        frame = ttk.Frame(parent); frame.grid(row=row, column=0, columnspan=2, sticky="ew", pady=2); frame.columnconfigure(1, weight=1)
        ttk.Label(frame, text=label + "：").grid(row=0, column=0, sticky="w")
        ttk.Entry(frame, textvariable=var).grid(row=0, column=1, sticky="ew", padx=4)
        cmd = (lambda: self._choose_dir(var)) if directory else (lambda: self._choose_file(var, [("JSON", "*.json")]))
        ttk.Button(frame, text="浏览…", command=cmd).grid(row=0, column=2)
        return frame

    def _version_visibility(self):
        if not hasattr(self, "version_rows"): return
        kind = self.version_source_type.get()
        for key, frame in self.version_rows.items():
            show = (kind == "json" and key == "json") or (kind == "excel" and key == "excel") or (kind == "txt" and key in {"origin", "translated"})
            if show: frame.grid()
            else: frame.grid_remove()

    def _load_version_baseline(self):
        kind = self.version_source_type.get()
        if kind == "json": return load_json_records(Path(self.version_json.get())).records
        if kind == "excel": return load_excel_records(Path(self.version_excel.get())).records
        origin = Path(self.version_origin.get()) if self.version_origin.get().strip() else Path(self.origin_dir.get())
        return load_txt_records(origin, Path(self.version_translated.get()), self.origin_encoding.get(), self.translated_encoding.get()).records

    def _do_version_compare(self):
        try:
            baseline = list(self._load_version_baseline())
            current = self.active_records()
            self.version_results = compare_translation_versions(current, baseline)
            self._clear_tree(self.version_tree); self.result_maps["versions"] = {}
            for i, x in enumerate(self.version_results):
                iid = f"ver{i}"; self.result_maps["versions"][iid] = x.record
                self.version_tree.insert("", "end", iid=iid, values=(x.record.file_key, self._display(x.record.original), self._display(x.baseline_translation),
                    self._display(x.current_translation), x.baseline_location, x.record.translated_location))
            self.status_var.set(f"原文一致且译文发生修改：{len(self.version_results)} 条。")
        except Exception as exc: messagebox.showerror("版本比较失败", str(exc), parent=self)

    def _export_version_report(self):
        if not self.version_results:
            messagebox.showinfo("版本报告", "请先比较版本。", parent=self); return
        path = filedialog.asksaveasfilename(parent=self, defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx"), ("CSV", "*.csv"), ("日志", "*.log")])
        if not path: return
        headers = ["文件", "原文", "旧译文", "新译文", "旧版位置", "新版位置"]
        rows = [[x.record.file_key, x.record.original, x.baseline_translation, x.current_translation, x.baseline_location, x.record.translated_location] for x in self.version_results]
        self._write_report(Path(path), headers, rows, "版本译文差异")
        messagebox.showinfo("导出完成", path, parent=self)

    # ------------------------------------------------------------------
    # Selectable batch QA report
    # ------------------------------------------------------------------
    def _default_report_checks(self) -> set[str]:
        cfg = load_config(self.config_path)
        saved = cfg.get("batch_report_checks")
        valid = set(BATCH_QA_CHECK_LABELS)
        if isinstance(saved, list):
            selected = {str(x) for x in saved if str(x) in valid}
            if selected:
                return selected
        # v4.2.1 defaults to all checks so speaker-name and duplicate-text
        # inconsistencies are included unless the user deliberately excludes them.
        return valid

    def _choose_report_checks(self) -> set[str] | None:
        win = tk.Toplevel(self)
        win.title("选择批量检查项目")
        win.transient(self); win.grab_set(); win.resizable(False, False)
        result: dict[str, set[str] | None] = {"value": None}
        selected = self._default_report_checks()
        vars_: dict[str, tk.BooleanVar] = {}

        ttk.Label(win, text="请选择本次报告需要执行的检查。未勾选项目不会扫描，也不会出现在报告中。",
                  foreground="#444").grid(row=0, column=0, columnspan=2, sticky="w", padx=12, pady=(10, 6))

        groups = [
            ("格式类检查", BATCH_QA_FORMAT_KEYS),
            ("内容一致性 / 人工校对", BATCH_QA_CONTENT_KEYS),
        ]
        for col, (title, keys) in enumerate(groups):
            frame = ttk.LabelFrame(win, text=title, padding=8)
            frame.grid(row=1, column=col, sticky="nsew", padx=(12 if col == 0 else 5, 12 if col == 1 else 5), pady=4)
            for r, key in enumerate(keys):
                var = tk.BooleanVar(value=key in selected); vars_[key] = var
                ttk.Checkbutton(frame, text=BATCH_QA_CHECK_LABELS[key], variable=var).grid(row=r, column=0, sticky="w", pady=2)

        btns = ttk.Frame(win); btns.grid(row=2, column=0, columnspan=2, sticky="ew", padx=12, pady=(6, 10))
        def set_keys(keys, value=True):
            keys = set(keys)
            for key, var in vars_.items():
                var.set(value if key in keys else (var.get() if value else False))
        ttk.Button(btns, text="全选", command=lambda: [v.set(True) for v in vars_.values()]).pack(side="left")
        ttk.Button(btns, text="全不选", command=lambda: [v.set(False) for v in vars_.values()]).pack(side="left", padx=4)
        ttk.Button(btns, text="仅格式类", command=lambda: [v.set(k in set(BATCH_QA_FORMAT_KEYS)) for k, v in vars_.items()]).pack(side="left", padx=4)
        ttk.Button(btns, text="仅内容一致性", command=lambda: [v.set(k in set(BATCH_QA_CONTENT_KEYS)) for k, v in vars_.items()]).pack(side="left", padx=4)

        def accept():
            chosen = {k for k, v in vars_.items() if v.get()}
            if not chosen:
                messagebox.showwarning("未选择检查项", "至少选择一项检查。", parent=win); return
            result["value"] = chosen
            cfg = load_config(self.config_path); cfg["batch_report_checks"] = sorted(chosen); save_config(self.config_path, cfg)
            win.destroy()
        def cancel():
            result["value"] = None; win.destroy()
        ttk.Button(btns, text="取消", command=cancel).pack(side="right")
        ttk.Button(btns, text="确定并继续", command=accept).pack(side="right", padx=4)
        win.protocol("WM_DELETE_WINDOW", cancel)
        win.update_idletasks()
        try:
            x = self.winfo_rootx() + max(20, (self.winfo_width() - win.winfo_reqwidth()) // 2)
            y = self.winfo_rooty() + max(20, (self.winfo_height() - win.winfo_reqheight()) // 2)
            win.geometry(f"+{x}+{y}")
        except Exception:
            pass
        self.wait_window(win)
        return result["value"]

    def _collect_format_report_rows(self, checks=None):
        chosen = set(checks or self._default_report_checks())
        issues = collect_batch_qa_report_issues(
            self.active_records(), chosen,
            float(self.face_limit.get()), float(self.narr_limit.get())
        )
        return [(x.issue_type, x.detail, x.record) for x in issues]

    def _export_format_report(self):
        if not self.all_records:
            messagebox.showinfo("批量校对报告", "请先检查数据源。", parent=self); return
        checks = self._choose_report_checks()
        if checks is None:
            return
        try:
            items = self._collect_format_report_rows(checks)
        except Exception as exc:
            messagebox.showerror("批量检查失败", str(exc), parent=self); return
        if not items:
            messagebox.showinfo("批量校对报告", "所选检查项目没有发现问题。", parent=self); return
        path = filedialog.asksaveasfilename(parent=self, defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx"), ("CSV", "*.csv"), ("日志", "*.log")])
        if not path: return
        headers = ["文件", "标记", "说话人ID", "问题类型", "问题说明", "原文", "译文", "原文位置", "译文位置"]
        rows = [[r.file_key, r.marker, r.speaker_id, typ, detail, r.original, r.translated, r.original_location, r.translated_location] for typ, detail, r in items]
        self._write_report(Path(path), headers, rows, "批量校对报告")
        labels = "、".join(BATCH_QA_CHECK_LABELS[k] for k in BATCH_QA_CHECK_LABELS if k in checks)
        messagebox.showinfo("批量校对报告", f"已执行 {len(checks)} 项检查，共输出 {len(rows)} 条问题。\n检查项：{labels}\n\n{path}", parent=self)

    @staticmethod
    def _write_report(path: Path, headers, rows, sheet_name: str):
        path.parent.mkdir(parents=True, exist_ok=True)
        if path.suffix.lower() == ".csv":
            with path.open("w", encoding="utf-8-sig", newline="") as fh:
                w = csv.writer(fh); w.writerow(headers); w.writerows(rows)
            return
        if path.suffix.lower() == ".log":
            with path.open("w", encoding="utf-8", newline="\n") as fh:
                fh.write("\t".join(headers) + "\n")
                for row in rows: fh.write("\t".join(str(x).replace("\t", "    ").replace("\r", "").replace("\n", "\\n") for x in row) + "\n")
            return
        wb = Workbook(); ws = wb.active; ws.title = sheet_name[:31]
        ws.append(list(headers))
        for row in rows: ws.append(list(row))
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="4472C4"); cell.alignment = Alignment(vertical="center")
        ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
        widths = [26, 12, 18, 20, 36, 55, 55, 45, 45]
        for i in range(1, len(headers)+1):
            ws.column_dimensions[chr(64+i) if i <= 26 else "A"].width = widths[i-1] if i <= len(widths) else 24
        for row in ws.iter_rows(min_row=2):
            for cell in row: cell.alignment = Alignment(vertical="top", wrap_text=True)
        wb.save(path); wb.close()

    # ------------------------------------------------------------------
    # Explicit event-row double click (avoids selection timing issues)
    # ------------------------------------------------------------------
    def _double_to_editor_event(self, key, tree, event):
        iid = tree.identify_row(event.y)
        obj = self.result_maps.get(key, {}).get(iid)
        rec = obj.record if hasattr(obj, "record") else obj
        if isinstance(rec, DataRecord): self._send_to_editor(rec); return "break"


def main():
    app = RPGMakerProofreadingApp(); app.mainloop()


if __name__ == "__main__": main()
